import io
from datetime import date

import numpy as np
import numpy_financial as npf
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

st.set_page_config(page_title="CRE Underwriting Calculator", page_icon="📊", layout="wide")

# ---------------------------------------------------------------------------
# Brand palette -- neutral base (grays / off-white / charcoal) with purposeful
# color-coding reserved for metrics that have a real underwriting threshold
# (DSCR, cash-on-cash, IRR, equity multiple). Everything else stays neutral
# so color always means something instead of just decorating the page.
# ---------------------------------------------------------------------------

INK = "1F2430"        # primary text
MUTED = "6B7280"      # secondary text
LINE = "E3E6EA"       # borders / dividers
ACCENT = "3B5BA3"     # single structural accent (headers, dividers, buttons)
ACCENT_LIGHT = "A9BAD9"  # lighter tint of the accent, used only to break out one-time sale proceeds on the cash flow chart
GOOD = "1E8E5A"
GOOD_BG = "E8F5EE"
WARN = "B7791F"
WARN_BG = "FDF3E0"
BAD = "C0392B"
BAD_BG = "FBEAE9"

st.markdown(
    f"""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

    html, body, [class*="css"],
    .stApp, .stApp p, .stApp label, .stApp li,
    .hero-eyebrow, .hero-title, .hero-caption, .section-title, .section-caption,
    .card-label, .card-value, .card-badge, .sidebar-byline {{
        font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif !important;
    }}
    [data-testid="stIconMaterial"] {{ font-family: 'Material Symbols Rounded' !important; }}
    #MainMenu, footer {{ visibility: hidden; }}
    .block-container {{ padding-top: 2.2rem; max-width: 1100px; }}

    .hero-eyebrow {{
        color: #{ACCENT}; font-weight: 700; font-size: 0.76rem; letter-spacing: 0.1em;
        text-transform: uppercase; margin-bottom: 0.4rem; white-space: nowrap;
    }}
    .hero-title {{ color: #{INK}; font-weight: 800; font-size: 2.05rem; line-height: 1.15; margin: 0 0 0.55rem 0; }}
    .hero-caption {{ color: #{MUTED}; font-size: 0.98rem; max-width: 740px; line-height: 1.55; }}
    .hero-divider {{ height: 3px; width: 56px; background: #{ACCENT}; border-radius: 2px; margin: 1.15rem 0 1.7rem 0; }}

    .section-label {{ display: flex; align-items: center; gap: 0.6rem; margin: 0.3rem 0 0.2rem 0; }}
    .section-bar {{ width: 5px; height: 21px; background: #{ACCENT}; border-radius: 2px; }}
    .section-title {{ font-size: 1.22rem; font-weight: 700; color: #{INK}; }}
    .section-caption {{ color: #{MUTED}; font-size: 0.92rem; margin: 0.25rem 0 1.1rem 1.35rem; max-width: 760px; line-height: 1.55; }}

    .card-grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(210px, 1fr)); gap: 0.85rem; margin-bottom: 1.5rem; }}
    .card {{ background: #FFFFFF; border: 1px solid #{LINE}; border-radius: 10px; padding: 0.95rem 1.1rem; border-left: 4px solid #C7CCD4; }}
    .card.status-good {{ border-left-color: #{GOOD}; }}
    .card.status-warn {{ border-left-color: #{WARN}; }}
    .card.status-bad {{ border-left-color: #{BAD}; }}
    .card-label {{ color: #{MUTED}; font-size: 0.71rem; font-weight: 600; letter-spacing: 0.04em; text-transform: uppercase; margin-bottom: 0.4rem; white-space: nowrap; }}
    .card-value {{ color: #{INK}; font-size: 1.5rem; font-weight: 700; line-height: 1.1; }}
    .card-badge {{ display: inline-block; margin-top: 0.45rem; font-size: 0.71rem; font-weight: 600; padding: 0.14rem 0.55rem; border-radius: 20px; }}
    .card-badge.status-good {{ background: #{GOOD_BG}; color: #{GOOD}; }}
    .card-badge.status-warn {{ background: #{WARN_BG}; color: #{WARN}; }}
    .card-badge.status-bad {{ background: #{BAD_BG}; color: #{BAD}; }}

    .stDownloadButton button {{ background: #{INK}; color: #fff; border: none; border-radius: 8px; font-weight: 600; padding: 0.55rem 1.15rem; }}
    .stDownloadButton button:hover {{ background: #{ACCENT}; color: #fff; }}

    .sidebar-byline {{
        color: #{MUTED}; font-size: 0.78rem; font-weight: 600; letter-spacing: 0.03em;
        text-transform: uppercase; margin: 0 0 1rem 0; padding-bottom: 0.85rem;
        border-bottom: 1px solid #{LINE};
    }}
    </style>
    """,
    unsafe_allow_html=True,
)


def hero(eyebrow, title, caption):
    st.markdown(
        f'<div class="hero-eyebrow">{eyebrow}</div>'
        f'<div class="hero-title">{title}</div>'
        f'<div class="hero-caption">{caption}</div>'
        f'<div class="hero-divider"></div>',
        unsafe_allow_html=True,
    )


def section(title, caption=None):
    st.markdown(
        f'<div class="section-label"><div class="section-bar"></div>'
        f'<div class="section-title">{title}</div></div>',
        unsafe_allow_html=True,
    )
    if caption:
        st.markdown(f'<div class="section-caption">{caption}</div>', unsafe_allow_html=True)


STATUS_TEXT = {"good": "Healthy", "warn": "Marginal", "bad": "Weak"}


def card_html(label, value, status=None, badge=None):
    cls = f"card status-{status}" if status else "card"
    badge_html = ""
    if status:
        badge_html = f'<div class="card-badge status-{status}">{badge or STATUS_TEXT[status]}</div>'
    return f'<div class="{cls}"><div class="card-label">{label}</div><div class="card-value">{value}</div>{badge_html}</div>'


def render_cards(cards):
    st.markdown('<div class="card-grid">' + "".join(card_html(**c) for c in cards) + "</div>", unsafe_allow_html=True)


def status_for(value, good_at, warn_at):
    """good_at / warn_at are lower-bound thresholds on a scale where higher is better."""
    if value >= good_at:
        return "good"
    if value >= warn_at:
        return "warn"
    return "bad"


def format_money(value):
    """Shorthand dollar formatting so KPI tiles never overflow -- $1.75M instead
    of $1,750,000, which was getting cut off with "..." in narrower windows."""
    if abs(value) >= 1_000_000:
        return f"${value / 1_000_000:.2f}M"
    if abs(value) >= 1_000:
        return f"${value / 1_000:.0f}K"
    return f"${value:,.0f}"


def format_psf(value):
    return f"${value:,.2f}"


# ---------------------------------------------------------------------------
# Amortization helpers
# ---------------------------------------------------------------------------

def annual_debt_service(loan_amount, annual_rate, amort_years):
    """Standard fixed-rate, monthly-amortizing mortgage payment, annualized."""
    if loan_amount <= 0:
        return 0.0
    monthly_rate = annual_rate / 12
    n_periods = amort_years * 12
    if monthly_rate == 0:
        return loan_amount / amort_years
    monthly_payment = loan_amount * monthly_rate / (1 - (1 + monthly_rate) ** (-n_periods))
    return monthly_payment * 12


def remaining_loan_balance(loan_amount, annual_rate, amort_years, years_elapsed):
    """Outstanding principal after `years_elapsed` years of amortization."""
    if loan_amount <= 0:
        return 0.0
    monthly_rate = annual_rate / 12
    n_periods = amort_years * 12
    k_periods = years_elapsed * 12
    if monthly_rate == 0:
        return max(loan_amount * (1 - years_elapsed / amort_years), 0.0)
    balance = loan_amount * (
        (1 + monthly_rate) ** n_periods - (1 + monthly_rate) ** k_periods
    ) / ((1 + monthly_rate) ** n_periods - 1)
    return max(balance, 0.0)


# ---------------------------------------------------------------------------
# Header
# ---------------------------------------------------------------------------

hero(
    "Commercial Real Estate",
    "Real Estate Deal Underwriting Calculator",
    "Enter a deal's basic numbers to see its cap rate, DSCR, cash-on-cash return, and a full "
    "multi-year levered IRR -- the same math behind a typical acquisition underwriting model.",
)

# ---------------------------------------------------------------------------
# Sidebar inputs
# ---------------------------------------------------------------------------

st.sidebar.markdown('<div class="sidebar-byline">Prepared by Diesel De Luz</div>', unsafe_allow_html=True)
st.sidebar.header("Deal Inputs")

purchase_price = st.sidebar.number_input("Purchase Price ($)", min_value=100_000, value=5_000_000, step=50_000)
square_footage = st.sidebar.number_input(
    "Square Footage (SF)", min_value=0, value=45_000, step=1_000,
    help="Optional -- enter the property's rentable square footage to see price, NOI, and exit value on a per-SF basis.",
)
year1_noi = st.sidebar.number_input("Year 1 NOI ($)", min_value=0, value=300_000, step=5_000)
noi_growth = st.sidebar.slider("Annual NOI Growth Rate (%)", 0.0, 8.0, 3.0, 0.25) / 100

st.sidebar.divider()
st.sidebar.subheader("Financing")
ltv = st.sidebar.slider("Loan-to-Value (%)", 0, 90, 65, 5) / 100
interest_rate = st.sidebar.slider("Interest Rate (%)", 2.0, 10.0, 6.5, 0.125) / 100
amort_years = st.sidebar.selectbox("Amortization (years)", [20, 25, 30], index=2)

st.sidebar.divider()
st.sidebar.subheader("Hold & Exit")
hold_period = st.sidebar.slider("Hold Period (years)", 1, 10, 5)
exit_cap_rate = st.sidebar.slider("Exit Cap Rate (%)", 3.0, 10.0, 5.5, 0.25) / 100
selling_costs_pct = st.sidebar.slider("Selling Costs (%)", 0.0, 5.0, 2.0, 0.25) / 100

# ---------------------------------------------------------------------------
# Core underwriting metrics
# ---------------------------------------------------------------------------

loan_amount = purchase_price * ltv
equity_invested = purchase_price - loan_amount
going_in_cap_rate = year1_noi / purchase_price if purchase_price else 0
debt_service = annual_debt_service(loan_amount, interest_rate, amort_years)
dscr = year1_noi / debt_service if debt_service else float("inf")
year1_cash_flow = year1_noi - debt_service
cash_on_cash = year1_cash_flow / equity_invested if equity_invested else 0

purchase_price_psf = purchase_price / square_footage if square_footage else None
noi_psf = year1_noi / square_footage if square_footage else None

# ---------------------------------------------------------------------------
# KPI tiles
# ---------------------------------------------------------------------------

section("Deal Snapshot")

with st.expander("New to underwriting? What this section means"):
    st.markdown(
        "This is the first gut check on a deal, before financing or hold-period assumptions "
        "come into play. **Going-in cap rate** tells you the return the property would throw "
        "off in year one if you paid all cash -- higher generally means more income relative "
        "to price. **Year 1 DSCR** asks a different question: once you add a loan, does the "
        "property's income cover the loan payment with room to spare? Lenders live and die by "
        "this number. **Year 1 cash-on-cash return** is the actual cash you'd pocket this year "
        "for every dollar you personally invested, after the mortgage payment. **Equity invested** "
        "is simply your out-of-pocket check -- purchase price minus whatever the bank is lending."
    )

dscr_status = status_for(dscr, 1.25, 1.00)
coc_status = status_for(cash_on_cash * 100, 8, 0)

render_cards([
    {"label": "Going-in Cap Rate", "value": f"{going_in_cap_rate * 100:.2f}%"},
    {"label": "Year 1 DSCR", "value": f"{dscr:.2f}x", "status": dscr_status},
    {"label": "Year 1 Cash-on-Cash", "value": f"{cash_on_cash * 100:.2f}%", "status": coc_status},
    {"label": "Equity Invested", "value": format_money(equity_invested)},
])

if dscr < 1.20:
    st.warning(
        f"DSCR of {dscr:.2f}x is below the 1.20x-1.25x minimum most lenders require -- this loan "
        "amount or rate combination likely wouldn't qualify as structured."
    )

# ---------------------------------------------------------------------------
# Multi-year proforma + levered IRR
# ---------------------------------------------------------------------------

years = list(range(1, hold_period + 1))
noi_by_year = [year1_noi * (1 + noi_growth) ** (y - 1) for y in years]
cash_flow_by_year = [noi - debt_service for noi in noi_by_year]

exit_year_noi = year1_noi * (1 + noi_growth) ** hold_period  # NOI in the year after sale
exit_value = exit_year_noi / exit_cap_rate if exit_cap_rate else 0
selling_costs = exit_value * selling_costs_pct
loan_payoff = remaining_loan_balance(loan_amount, interest_rate, amort_years, hold_period)
net_sale_proceeds = exit_value - selling_costs - loan_payoff

proforma = pd.DataFrame({
    "Year": years,
    "NOI": noi_by_year,
    "Debt Service": [debt_service] * hold_period,
    "Cash Flow": cash_flow_by_year,
})
proforma.loc[proforma["Year"] == hold_period, "Cash Flow"] += net_sale_proceeds

levered_cash_flows = [-equity_invested] + proforma["Cash Flow"].tolist()
irr = npf.irr(levered_cash_flows)
total_distributions = sum(proforma["Cash Flow"])
equity_multiple = total_distributions / equity_invested if equity_invested else 0

exit_value_psf = exit_value / square_footage if square_footage else None

st.divider()
section(
    "Hold-Period Returns",
    "Full multi-year cash flow projection plus the levered IRR and equity multiple across your assumed hold period.",
)

with st.expander("New to underwriting? What this section means"):
    st.markdown(
        "Deals aren't judged on year one alone -- most investors buy planning to hold for "
        "several years and then sell. This section runs the whole holding period forward: NOI "
        "grows each year, cash flow builds after the mortgage payment, and in the final year the "
        "model adds what you'd walk away with from selling the property (its exit value, minus "
        "selling costs and whatever's left on the loan). **Levered IRR** is the annualized return "
        "across that entire timeline, accounting for when money went in and when it came back. "
        "**Equity multiple** is simpler: total dollars back divided by dollars invested -- a 2.0x "
        "means you doubled your money over the hold period."
    )

irr_pct = irr * 100 if irr is not None else None
irr_status = status_for(irr_pct, 15, 8) if irr_pct is not None else None
multiple_status = status_for(equity_multiple, 1.8, 1.2)

return_cards = [
    {"label": "Levered IRR", "value": f"{irr_pct:.1f}%" if irr_pct is not None else "n/a", "status": irr_status},
    {"label": "Equity Multiple", "value": f"{equity_multiple:.2f}x", "status": multiple_status},
    {"label": "Exit Value", "value": format_money(exit_value)},
]
render_cards(return_cards)

st.dataframe(
    proforma,
    column_config={
        "NOI": st.column_config.NumberColumn(format="$%d"),
        "Debt Service": st.column_config.NumberColumn(format="$%d"),
        "Cash Flow": st.column_config.NumberColumn(format="$%d"),
    },
    hide_index=True,
    width="stretch",
)

cash_flow_fig = go.Figure()
cash_flow_fig.add_bar(
    x=years, y=cash_flow_by_year, name="Operating Cash Flow",
    marker_color=f"#{ACCENT}", yaxis="y1",
    hovertemplate="Year %{x}<br>Operating: %{y:$,.0f}<extra></extra>",
)
cash_flow_fig.add_bar(
    x=[years[-1]], y=[net_sale_proceeds], name="Net Sale Proceeds (exit year)",
    marker_color=f"#{ACCENT_LIGHT}", yaxis="y2",
    hovertemplate="Year %{x}<br>Sale Proceeds: %{y:$,.0f}<extra></extra>",
)
cash_flow_fig.update_layout(
    barmode="group",
    bargap=0.3,
    template="plotly_white",
    font=dict(family="Inter, -apple-system, sans-serif", color=f"#{INK}", size=13),
    xaxis=dict(title="Year", tickmode="linear", dtick=1, showgrid=False),
    yaxis=dict(
        title=dict(text="Operating Cash Flow", font=dict(color=f"#{ACCENT}", size=12)),
        tickfont=dict(color=f"#{ACCENT}"),
        showgrid=True, gridcolor=f"#{LINE}", tickformat="$,.0f", zeroline=False,
    ),
    yaxis2=dict(
        title=dict(text="Sale Proceeds", font=dict(color=f"#{ACCENT_LIGHT}", size=12)),
        tickfont=dict(color=f"#{ACCENT_LIGHT}"),
        overlaying="y", side="right", showgrid=False, tickformat="$,.0f", zeroline=False,
    ),
    legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1, title=None),
    margin=dict(l=10, r=10, t=10, b=10),
    height=340,
    hovermode="x unified",
    plot_bgcolor="white",
    paper_bgcolor="white",
)
st.plotly_chart(cash_flow_fig, config={"displayModeBar": False})
st.caption(
    "Two separate scales -- operating cash flow (left axis) is measured in the "
    "tens of thousands, while the one-time sale proceeds (right axis) run into "
    "the millions. Splitting the axes keeps Years 1-4 readable instead of "
    "flattening them next to the exit-year payout."
)

# ---------------------------------------------------------------------------
# Per square foot
# ---------------------------------------------------------------------------

st.divider()
section(
    "Per Square Foot",
    "The same deal viewed on a price-per-square-foot basis -- the way most brokers and appraisers quote a market comp.",
)

with st.expander("New to underwriting? What this section means"):
    st.markdown(
        "Purchase price alone doesn't tell you if a deal is cheap or expensive -- a $5M property "
        "could be a great deal or a bad one depending on its size. Real estate professionals "
        "normalize for this by dividing key numbers by square footage, the same way home shoppers "
        "compare price per square foot. It lets you compare this deal against other listings or "
        "recent sales on an apples-to-apples basis, no matter how big either property is."
    )

if square_footage:
    render_cards([
        {"label": "Purchase Price / SF", "value": format_psf(purchase_price_psf)},
        {"label": "Year 1 NOI / SF", "value": format_psf(noi_psf)},
        {"label": "Exit Value / SF", "value": format_psf(exit_value_psf)},
    ])
else:
    st.info("Enter a square footage in the sidebar to see price, NOI, and exit value on a per-SF basis.")

# ---------------------------------------------------------------------------
# Sensitivity analysis: IRR across exit cap rate x NOI growth rate
# ---------------------------------------------------------------------------

st.divider()
section(
    "Sensitivity Analysis: Levered IRR",
    "How the deal's IRR shifts under different exit cap rate and NOI growth assumptions -- the same "
    "kind of sensitivity table used to stress-test a deal before committing capital.",
)


with st.expander("New to underwriting? What this section means"):
    st.markdown(
        "Every underwriting model rests on assumptions about the future -- and the two hardest "
        "to predict are what cap rate the market will demand when you sell, and how fast rents "
        "will grow along the way. This table reruns the entire IRR calculation across a range of "
        "both, so instead of one single return number, you can see how the deal performs across "
        "a more optimistic and a more pessimistic future. A deal that still looks solid across "
        "most of this grid is more resilient than one that only works if everything goes right."
    )


def irr_for(exit_cap, growth):
    noi_by_yr = [year1_noi * (1 + growth) ** (y - 1) for y in years]
    cf_by_yr = [noi - debt_service for noi in noi_by_yr]
    exit_noi = year1_noi * (1 + growth) ** hold_period
    ev = exit_noi / exit_cap if exit_cap else 0
    sc = ev * selling_costs_pct
    payoff = remaining_loan_balance(loan_amount, interest_rate, amort_years, hold_period)
    proceeds = ev - sc - payoff
    cfs = cf_by_yr.copy()
    cfs[-1] += proceeds
    return npf.irr([-equity_invested] + cfs)


exit_cap_grid = [exit_cap_rate + delta for delta in (-0.01, -0.005, 0, 0.005, 0.01)]
growth_grid = [noi_growth + delta for delta in (-0.02, -0.01, 0, 0.01, 0.02)]

sensitivity = pd.DataFrame(
    {f"{g * 100:.1f}% growth": [irr_for(ec, g) for ec in exit_cap_grid] for g in growth_grid},
    index=[f"{ec * 100:.2f}% exit cap" for ec in exit_cap_grid],
)

st.dataframe(
    sensitivity.style.format("{:.1%}").background_gradient(cmap="RdYlGn", axis=None),
    width="stretch",
)

st.divider()
with st.expander("How these numbers are calculated"):
    st.markdown(
        "**Going-in cap rate** = Year 1 NOI ÷ Purchase Price.\n\n"
        "**Debt service** uses a standard fixed-rate, monthly-amortizing loan payment, annualized.\n\n"
        "**DSCR** = NOI ÷ annual debt service -- most lenders require at least 1.20x-1.25x.\n\n"
        "**Cash-on-cash return** = (NOI − debt service) ÷ equity invested, for the first year.\n\n"
        "**Exit value** = the NOI in the year *after* the sale, divided by the exit cap rate (the "
        "standard convention, since a buyer is underwriting the property's *next* year of income).\n\n"
        "**Levered IRR** discounts the initial equity outflow against every year's cash flow plus "
        "net sale proceeds (exit value, minus selling costs, minus the remaining loan balance) in "
        "the final year.\n\n"
        "**Color coding** on DSCR, cash-on-cash, IRR, and equity multiple reflects common underwriting "
        "thresholds (e.g. DSCR under 1.0x is red because the property can't cover its own debt service). "
        "Going-in cap rate and per-SF figures aren't color-coded -- what counts as \"good\" there depends "
        "entirely on market and asset class, so a fixed threshold would be misleading."
    )

# ---------------------------------------------------------------------------
# Export: formatted Excel workbook with live formulas
# ---------------------------------------------------------------------------


def build_excel_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Underwriting Model"

    header_fill = PatternFill("solid", fgColor=INK)
    header_font = Font(color="FFFFFF", bold=True, size=12)
    label_font = Font(bold=True, color=INK)
    section_fill = PatternFill("solid", fgColor="F2F3F5")
    thin = Side(style="thin", color="D9DCE1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    money_fmt = '$#,##0'
    money2_fmt = '$#,##0.00'
    pct_fmt = '0.00%'
    mult_fmt = '0.00"x"'

    def set_header(cell_range, text):
        ws.merge_cells(cell_range)
        c = ws[cell_range.split(":")[0]]
        c.value = text
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="left", vertical="center")

    def label_cell(row, text):
        c = ws.cell(row=row, column=1, value=text)
        c.font = label_font
        c.fill = section_fill

    def value_cell(row, col, value, number_format=None, formula=False):
        c = ws.cell(row=row, column=col, value=value)
        c.border = border
        if number_format:
            c.number_format = number_format
        return c

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 15

    set_header("A1:B1", "Deal Inputs")
    inputs = [
        ("Purchase Price ($)", purchase_price, money_fmt),
        ("Square Footage (SF)", square_footage if square_footage else 0, "#,##0"),
        ("Year 1 NOI ($)", year1_noi, money_fmt),
        ("Annual NOI Growth Rate", noi_growth, pct_fmt),
        ("Loan-to-Value", ltv, pct_fmt),
        ("Interest Rate", interest_rate, pct_fmt),
        ("Amortization (years)", amort_years, "0"),
        ("Hold Period (years)", hold_period, "0"),
        ("Exit Cap Rate", exit_cap_rate, pct_fmt),
        ("Selling Costs", selling_costs_pct, pct_fmt),
    ]
    row = 2
    input_rows = {}
    for label, val, fmt in inputs:
        label_cell(row, label)
        value_cell(row, 2, val, fmt)
        input_rows[label] = row
        row += 1

    r_price = input_rows["Purchase Price ($)"]
    r_sf = input_rows["Square Footage (SF)"]
    r_noi = input_rows["Year 1 NOI ($)"]
    r_growth = input_rows["Annual NOI Growth Rate"]
    r_ltv = input_rows["Loan-to-Value"]
    r_rate = input_rows["Interest Rate"]
    r_amort = input_rows["Amortization (years)"]
    r_hold = input_rows["Hold Period (years)"]
    r_exitcap = input_rows["Exit Cap Rate"]
    r_sellcost = input_rows["Selling Costs"]

    row += 1
    set_header(f"A{row}:B{row}", "Key Metrics (live formulas)")
    row += 1
    r_loan = row
    label_cell(row, "Loan Amount")
    value_cell(row, 2, f"=B{r_price}*B{r_ltv}", money_fmt)
    row += 1
    r_equity = row
    label_cell(row, "Equity Invested")
    value_cell(row, 2, f"=B{r_price}-B{r_loan}", money_fmt)
    row += 1
    r_cap = row
    label_cell(row, "Going-in Cap Rate")
    value_cell(row, 2, f"=B{r_noi}/B{r_price}", pct_fmt)
    row += 1
    r_ds = row
    label_cell(row, "Annual Debt Service")
    value_cell(row, 2, f"=-PMT(B{r_rate}/12,B{r_amort}*12,B{r_loan})*12", money_fmt)
    row += 1
    r_dscr = row
    label_cell(row, "Year 1 DSCR")
    value_cell(row, 2, f"=B{r_noi}/B{r_ds}", mult_fmt)
    row += 1
    r_cf1 = row
    label_cell(row, "Year 1 Cash Flow")
    value_cell(row, 2, f"=B{r_noi}-B{r_ds}", money_fmt)
    row += 1
    r_coc = row
    label_cell(row, "Year 1 Cash-on-Cash")
    value_cell(row, 2, f"=B{r_cf1}/B{r_equity}", pct_fmt)
    row += 1
    r_price_psf = row
    label_cell(row, "Purchase Price / SF")
    value_cell(row, 2, f"=IF(B{r_sf}=0,\"\",B{r_price}/B{r_sf})", money2_fmt)
    row += 1
    r_noi_psf = row
    label_cell(row, "Year 1 NOI / SF")
    value_cell(row, 2, f"=IF(B{r_sf}=0,\"\",B{r_noi}/B{r_sf})", money2_fmt)

    row += 2
    set_header(f"A{row}:E{row}", "Hold-Period Proforma")
    row += 1
    hdr_row = row
    for i, h in enumerate(["Year", "NOI", "Debt Service", "Cash Flow"]):
        c = ws.cell(row=hdr_row, column=1 + i, value=h)
        c.font = Font(bold=True, color=INK)
        c.fill = section_fill
        c.border = border
    row += 1
    r_year0 = row
    ws.cell(row=row, column=1, value=0).border = border
    value_cell(row, 2, None, money_fmt)
    value_cell(row, 3, None, money_fmt)
    value_cell(row, 4, f"=-B{r_equity}", money_fmt)
    row += 1
    first_data_row = row
    for i, y in enumerate(years):
        c_year = ws.cell(row=row, column=1, value=y)
        c_year.border = border
        if i == 0:
            noi_formula = f"=B{r_noi}"
        else:
            noi_formula = f"=B{row - 1}*(1+B{r_growth})"
        value_cell(row, 2, noi_formula, money_fmt)
        value_cell(row, 3, f"=$B${r_ds}", money_fmt)
        value_cell(row, 4, f"=B{row}-C{row}", money_fmt)
        row += 1
    last_data_row = row - 1

    row += 1
    r_exit_noi = row
    label_cell(row, "Exit Year NOI (year after sale)")
    value_cell(row, 2, f"=B{last_data_row}*(1+B{r_growth})", money_fmt)
    row += 1
    r_exit_val = row
    label_cell(row, "Exit Value")
    value_cell(row, 2, f"=B{r_exit_noi}/B{r_exitcap}", money_fmt)
    row += 1
    r_exit_psf = row
    label_cell(row, "Exit Value / SF")
    value_cell(row, 2, f"=IF(B{r_sf}=0,\"\",B{r_exit_val}/B{r_sf})", money2_fmt)
    row += 1
    r_sellcosts = row
    label_cell(row, "Selling Costs")
    value_cell(row, 2, f"=B{r_exit_val}*B{r_sellcost}", money_fmt)
    row += 1
    r_payoff = row
    label_cell(row, "Loan Payoff at Exit")
    value_cell(
        row, 2,
        f"=B{r_loan}*(((1+B{r_rate}/12)^(B{r_amort}*12))-((1+B{r_rate}/12)^(B{r_hold}*12)))"
        f"/(((1+B{r_rate}/12)^(B{r_amort}*12))-1)",
        money_fmt,
    )
    row += 1
    r_proceeds = row
    label_cell(row, "Net Sale Proceeds")
    value_cell(row, 2, f"=B{r_exit_val}-B{r_sellcosts}-B{r_payoff}", money_fmt)

    # Fold net sale proceeds into the final proforma year's cash flow
    ws.cell(row=last_data_row, column=4).value = f"=B{last_data_row}-C{last_data_row}+B{r_proceeds}"

    row += 2
    set_header(f"A{row}:B{row}", "Levered Returns")
    row += 1
    r_irr = row
    label_cell(row, "Levered IRR")
    irr_range = f"D{r_year0}:D{last_data_row}"
    value_cell(row, 2, f"=IRR({irr_range})", pct_fmt)
    row += 1
    r_multiple = row
    label_cell(row, "Equity Multiple")
    cf_range = f"D{first_data_row}:D{last_data_row}"
    value_cell(row, 2, f"=SUM({cf_range})/B{r_equity}", mult_fmt)

    ws.freeze_panes = "A2"
    return wb


st.divider()
section("Export", "Download this deal as a live spreadsheet model or a formatted summary PDF.")

col_xlsx, col_pdf = st.columns(2)

with col_xlsx:
    wb = build_excel_workbook()
    buf = io.BytesIO()
    wb.save(buf)
    st.download_button(
        "Download Excel Model (.xlsx)",
        data=buf.getvalue(),
        file_name="cre_underwriting_model.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch",
    )
    st.caption("Every downstream cell is a live formula -- change an input in Excel and the model recalculates.")


def build_pdf_report():
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=letter,
        topMargin=0.6 * inch, bottomMargin=0.6 * inch, leftMargin=0.65 * inch, rightMargin=0.65 * inch,
    )
    styles = {
        "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=18, leading=23, textColor=colors.HexColor(f"#{INK}"), spaceAfter=6),
        "sub": ParagraphStyle("sub", fontName="Helvetica", fontSize=9.5, leading=13, textColor=colors.HexColor(f"#{MUTED}"), spaceAfter=14),
        "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12.5, leading=16, textColor=colors.HexColor(f"#{INK}"), spaceBefore=14, spaceAfter=6),
    }
    story = [
        Paragraph("Real Estate Deal Underwriting Summary", styles["title"]),
        Paragraph(f"Generated {date.today().strftime('%B %d, %Y')}", styles["sub"]),
    ]

    def status_color(status):
        return {"good": GOOD, "warn": WARN, "bad": BAD, None: INK}[status]

    story.append(Paragraph("Deal Assumptions", styles["h2"]))
    sf_line = f"{square_footage:,.0f} SF" if square_footage else "Not provided"
    assumptions = [
        ["Purchase Price", f"${purchase_price:,.0f}", "Square Footage", sf_line],
        ["Year 1 NOI", f"${year1_noi:,.0f}", "NOI Growth Rate", f"{noi_growth * 100:.2f}%"],
        ["Loan-to-Value", f"{ltv * 100:.0f}%", "Interest Rate", f"{interest_rate * 100:.2f}%"],
        ["Amortization", f"{amort_years} yrs", "Hold Period", f"{hold_period} yrs"],
        ["Exit Cap Rate", f"{exit_cap_rate * 100:.2f}%", "Selling Costs", f"{selling_costs_pct * 100:.2f}%"],
    ]
    t = Table(assumptions, colWidths=[1.55 * inch, 1.55 * inch, 1.55 * inch, 1.55 * inch])
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor(f"#{INK}")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{LINE}")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F3F5")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F2F3F5")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(t)

    story.append(Paragraph("Key Metrics", styles["h2"]))
    metric_rows = [
        ["Going-in Cap Rate", f"{going_in_cap_rate * 100:.2f}%", None],
        ["Year 1 DSCR", f"{dscr:.2f}x", dscr_status],
        ["Year 1 Cash-on-Cash", f"{cash_on_cash * 100:.2f}%", coc_status],
        ["Equity Invested", format_money(equity_invested), None],
        ["Levered IRR", f"{irr_pct:.1f}%" if irr_pct is not None else "n/a", irr_status],
        ["Equity Multiple", f"{equity_multiple:.2f}x", multiple_status],
        ["Exit Value", format_money(exit_value), None],
    ]
    table_data = [["Metric", "Value"]] + [[r[0], r[1]] for r in metric_rows]
    t2 = Table(table_data, colWidths=[3.1 * inch, 3.1 * inch])
    style_cmds = [
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{INK}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{LINE}")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for i, r in enumerate(metric_rows, start=1):
        if r[2]:
            style_cmds.append(("TEXTCOLOR", (1, i), (1, i), colors.HexColor(f"#{status_color(r[2])}")))
            style_cmds.append(("FONTNAME", (1, i), (1, i), "Helvetica-Bold"))
    t2.setStyle(TableStyle(style_cmds))
    story.append(t2)

    if square_footage:
        story.append(Paragraph("Per Square Foot", styles["h2"]))
        psf_rows = [
            ["Metric", "Value"],
            ["Purchase Price / SF", format_psf(purchase_price_psf)],
            ["Year 1 NOI / SF", format_psf(noi_psf)],
            ["Exit Value / SF", format_psf(exit_value_psf)],
        ]
        t3 = Table(psf_rows, colWidths=[3.1 * inch, 3.1 * inch])
        t3.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 9.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F3F5")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{LINE}")),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(t3)

    story.append(Paragraph("Hold-Period Proforma", styles["h2"]))
    proforma_rows = [["Year", "NOI", "Debt Service", "Cash Flow"]] + [
        [
            str(int(row["Year"])),
            f"${row['NOI']:,.0f}",
            f"${row['Debt Service']:,.0f}",
            f"${row['Cash Flow']:,.0f}",
        ]
        for _, row in proforma.iterrows()
    ]
    t4 = Table(proforma_rows, colWidths=[1.55 * inch] * 4)
    t4.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{INK}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{LINE}")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(t4)
    story.append(Spacer(1, 10))
    story.append(Paragraph(
        "Generated by the Real Estate Deal Underwriting Calculator -- re-underwriting-calculator.streamlit.app",
        styles["sub"],
    ))

    doc.build(story)
    return buf.getvalue()


with col_pdf:
    pdf_bytes = build_pdf_report()
    st.download_button(
        "Download PDF Summary",
        data=pdf_bytes,
        file_name="cre_underwriting_summary.pdf",
        mime="application/pdf",
        width="stretch",
    )
    st.caption("A clean, formatted one-page summary of the deal assumptions, key metrics, and proforma.")
